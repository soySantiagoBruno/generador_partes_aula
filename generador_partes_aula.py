#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GENERADOR AUTOMÁTICO DE PARTES DE AULA CON INTERFAZ GRÁFICA (VERSIÓN EXE)
"""

import os
import sys
import copy
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PLANTILLA_DEFAULT = "plantilla_base.docx"

MESES_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
    5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
    9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
}

DATOS_POR_DEFECTO = {
    "config": {
        "LUGAR": "CAMPO DE MAYO",
        "DIA": "      ", 
        "MES": MESES_ES[datetime.now().month],
        "ANIO": str(datetime.now().year),
        "TURNO": "MAÑANA",
        "GENERO": "MASCULINO",
        "OFICIAL_SEMANA": "Teniente Martin GONZALEZ POSSI",
        "SUBOFICIAL_SEMANA": "Cabo Primero Alejandro FERNANDEZ",
        "ARCHIVO_SALIDA": "Parte_Aula_Generado.docx"
    },
    "secciones": {
        "CAMPAÑA / RADAR": [
            {"grado": "AP Ier Año", "esp": "MEC EQ CAMP", "nombre": "CALISAYA LEANDRO IVAN", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC EQ CAMP", "nombre": "MALDONADO RAMIRO SAMUEL", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC RAR", "nombre": "LEAÑO MATIAS JOSÉ", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC RAR", "nombre": "LESCANO MIRANDA YONATAN", "estado": "P", "obs": ""},
        ],
        "INFORMATICA": [
            {"grado": "AP Ier Año", "esp": "MEC INFO", "nombre": "BRUNO SANTIAGO", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC INFO", "nombre": "GONZALES LEONARDO JOEL", "estado": "P", "obs": ""},
        ],
        "MEC EQ FIJ": [
            {"grado": "AP Ier Año", "esp": "MEC EQ FIJ", "nombre": "BONAZZOLA MAURICIO", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC EQ FIJ", "nombre": "CASANOVA PATRICIO", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC EQ FIJ", "nombre": "FERREIRA MARCOS ADRIAN", "estado": "P", "obs": ""},
            {"grado": "AP Ier Año", "esp": "MEC EQ FIJ", "nombre": "VERA TOBIAS AGUSTIN", "estado": "P", "obs": ""},
        ]
    }
}

# =============================================================================
# FUNCIONES DE RUTAS PARA EL EXE
# =============================================================================

def get_resource_path(relative_path):
    """Obtiene la ruta a los archivos empaquetados dentro del .exe"""
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def get_output_dir():
    """Obtiene la carpeta donde el usuario tiene guardado el .exe"""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.abspath(os.path.dirname(__file__))

# =============================================================================
# FUNCIONES DE MANEJO DE WORD
# =============================================================================

def cargar_desde_excel(ruta_excel):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        config = copy.deepcopy(DATOS_POR_DEFECTO["config"])
        if "Configuracion" in wb.sheetnames:
            ws_cfg = wb["Configuracion"]
            for row in ws_cfg.iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    clave = str(row[0]).strip().upper()
                    valor = "" if row[1] is None else str(row[1]).strip()
                    config[clave] = valor

        secciones = {"CAMPAÑA / RADAR": [], "INFORMATICA": [], "MEC EQ FIJ": []}
        if "Personal" in wb.sheetnames:
            ws_per = wb["Personal"]
            for row in ws_per.iter_rows(min_row=2, values_only=True):
                if not row or not row[0] or not row[3]: continue
                sec = str(row[0]).strip().upper()
                grado = str(row[1] or "AP Ier Año").strip()
                esp = str(row[2] or "").strip()
                nombre = str(row[3] or "").strip()
                estado = str(row[4] or "P").strip().upper()
                obs = "" if row[5] is None else str(row[5]).strip()
                if sec not in secciones:
                    secciones[sec] = []
                secciones[sec].append({
                    "grado": grado, "esp": esp, "nombre": nombre,
                    "estado": "A" if estado == "A" else "P", "obs": obs
                })
        return {"config": config, "secciones": secciones}
    except Exception:
        return copy.deepcopy(DATOS_POR_DEFECTO)

def actualizar_linea_fe_p_a(parrafo, fe, p, a):
    for run in parrafo.runs:
        run.text = ""
    espacios = " " * 59
    partes = [
        ("FE", True), (f": {fe:02d}{espacios}", False),
        ("P", True), (f": {p:02d}{espacios}", False),
        ("A", True), (f": {a:02d}", False)
    ]
    for texto, negrita in partes:
        r = parrafo.add_run(texto)
        r.bold = negrita
        r.font.size = Pt(12)

def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for border_name in ['top', 'left', 'bottom', 'right']:
        element = tcBorders.find(qn(f'w:{border_name}'))
        if element is None:
            element = OxmlElement(f'w:{border_name}')
            tcBorders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')

def rellenar_tabla_seccion(tabla, alumnos):
    while len(tabla.rows) > 1:
        tbl = tabla._tbl
        tbl.remove(tabla.rows[-1]._tr)
    for idx, alum in enumerate(alumnos, start=1):
        row_cells = tabla.add_row().cells
        es_presente = (alum["estado"].upper() != "A")
        valores = [
            f"{idx:02d}", alum["grado"], alum["esp"], alum["nombre"],
            "X" if es_presente else "", "" if es_presente else "X", alum["obs"]
        ]
        for c_idx, val in enumerate(valores):
            cell = row_cells[c_idx]
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.text = ""
            if c_idx in [0, 1, 2, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(11)

def reemplazar_en_cuadros_de_texto(doc, cfg):
    xml_tree = doc.element
    txbx_paragraphs = xml_tree.xpath('//w:txbxContent//w:p')
    dia_str = cfg.get("DIA", "      ")
    if not dia_str.strip():
        dia_str = "      "
    texto_fecha = f"{cfg['LUGAR']}, {dia_str} de {cfg['MES']} de {cfg['ANIO']}."

    for p in txbx_paragraphs:
        texto_completo = "".join(p.xpath('.//w:t/text()'))
        runs = p.xpath('.//w:r')
        if not runs: continue
        if "CAMPO DE MAYO" in texto_completo or "de 202" in texto_completo:
            t_nodes = p.xpath('.//w:t')
            if t_nodes:
                t_nodes[0].text = texto_fecha
                for t_extra in t_nodes[1:]: t_extra.text = ""
        elif "GONZALEZ POSSI" in texto_completo:
            t_nodes = p.xpath('.//w:t')
            if t_nodes:
                t_nodes[0].text = cfg["OFICIAL_SEMANA"]
                for t_extra in t_nodes[1:]: t_extra.text = ""
        elif "Alejandro FERNANDEZ" in texto_completo:
            t_nodes = p.xpath('.//w:t')
            if t_nodes:
                t_nodes[0].text = cfg["SUBOFICIAL_SEMANA"]
                for t_extra in t_nodes[1:]: t_extra.text = ""

def generar_parte_word(datos, ruta_salida):
    # Busca la plantilla empaquetada dentro del .exe
    ruta_plantilla = get_resource_path(PLANTILLA_DEFAULT)
    
    if not os.path.exists(ruta_plantilla):
        raise FileNotFoundError("Error crítico: No se encontró la plantilla empaquetada.")

    doc = docx.Document(ruta_plantilla)
    cfg = datos["config"]
    secciones = datos["secciones"]

    turno = cfg.get("TURNO", "MAÑANA").upper()
    genero = cfg.get("GENERO", "MASCULINO").upper()

    mapa_parrafos_fe = {"CAMPAÑA / RADAR": 8, "INFORMATICA": 11, "MEC EQ FIJ": 14}
    mapa_tablas = {"CAMPAÑA / RADAR": 0, "INFORMATICA": 1, "MEC EQ FIJ": 2}

    for idx_p in [4, 7, 10, 13]:
        if idx_p < len(doc.paragraphs):
            p = doc.paragraphs[idx_p]
            nuevo_txt = p.text.replace("MAÑANA", turno).replace("MASCULINO", genero)
            if nuevo_txt != p.text and p.runs:
                p.runs[0].text = nuevo_txt
                for r in p.runs[1:]: r.text = ""

    resumen_total = {"FE": 0, "P": 0, "A": 0}
    for nombre_sec, idx_tabla in mapa_tablas.items():
        alumnos = secciones.get(nombre_sec, [])
        fe = len(alumnos)
        p_cnt = sum(1 for a in alumnos if a["estado"].upper() != "A")
        a_cnt = fe - p_cnt

        resumen_total["FE"] += fe
        resumen_total["P"] += p_cnt
        resumen_total["A"] += a_cnt

        idx_fe = mapa_parrafos_fe[nombre_sec]
        actualizar_linea_fe_p_a(doc.paragraphs[idx_fe], fe, p_cnt, a_cnt)
        rellenar_tabla_seccion(doc.tables[idx_tabla], alumnos)

    reemplazar_en_cuadros_de_texto(doc, cfg)
    
    # Guarda el archivo en la misma carpeta donde el usuario tiene el .exe
    ruta_salida_final = os.path.join(get_output_dir(), ruta_salida)
    doc.save(ruta_salida_final)
    
    return ruta_salida_final, resumen_total

# =============================================================================
# INTERFAZ GRÁFICA (GUI)
# =============================================================================

class AppPartes:
    def __init__(self, root, datos):
        self.root = root
        self.root.title("Generador de Partes de Aula - 3ra Ca Ars")
        self.root.geometry("800x600")
        self.datos = datos
        
        self.mes_var = tk.StringVar(value=self.datos["config"]["MES"])
        self.turno_var = tk.StringVar(value=self.datos["config"]["TURNO"])
        self.oficial_var = tk.StringVar(value=self.datos["config"]["OFICIAL_SEMANA"])
        self.suboficial_var = tk.StringVar(value=self.datos["config"]["SUBOFICIAL_SEMANA"])
        self.asistencia_vars = {}

        self.crear_interfaz()

    def crear_interfaz(self):
        frame_top = ttk.LabelFrame(self.root, text="Configuración General", padding=(10, 5))
        frame_top.pack(fill=tk.X, padx=10, pady=5)

        ttk.Label(frame_top, text="Mes:").grid(row=0, column=0, sticky=tk.W, padx=5, pady=2)
        ttk.Entry(frame_top, textvariable=self.mes_var, width=15).grid(row=0, column=1, sticky=tk.W, padx=5, pady=2)

        ttk.Label(frame_top, text="Turno:").grid(row=0, column=2, sticky=tk.W, padx=5, pady=2)
        ttk.Combobox(frame_top, textvariable=self.turno_var, values=["MAÑANA", "TARDE"], state="readonly", width=12).grid(row=0, column=3, sticky=tk.W, padx=5, pady=2)

        ttk.Label(frame_top, text="Oficial Semana:").grid(row=1, column=0, sticky=tk.W, padx=5, pady=2)
        ttk.Entry(frame_top, textvariable=self.oficial_var, width=35).grid(row=1, column=1, columnspan=3, sticky=tk.W, padx=5, pady=2)

        ttk.Label(frame_top, text="Suboficial Semana:").grid(row=2, column=0, sticky=tk.W, padx=5, pady=2)
        ttk.Entry(frame_top, textvariable=self.suboficial_var, width=35).grid(row=2, column=1, columnspan=3, sticky=tk.W, padx=5, pady=2)

        notebook = ttk.Notebook(self.root)
        notebook.pack(expand=True, fill=tk.BOTH, padx=10, pady=5)

        for sec_nombre, alumnos in self.datos["secciones"].items():
            frame_tab = ttk.Frame(notebook, padding=10)
            notebook.add(frame_tab, text=sec_nombre)
            self.asistencia_vars[sec_nombre] = []
            
            ttk.Label(frame_tab, text="Estado", font=("Helvetica", 9, "bold")).grid(row=0, column=0, padx=5, pady=5)
            ttk.Label(frame_tab, text="Aspirante", font=("Helvetica", 9, "bold")).grid(row=0, column=1, sticky=tk.W, padx=5, pady=5)
            ttk.Label(frame_tab, text="Observación (Motivo de Ausencia)", font=("Helvetica", 9, "bold")).grid(row=0, column=2, sticky=tk.W, padx=5, pady=5)

            for i, alum in enumerate(alumnos):
                var_p = tk.IntVar(value=1 if alum["estado"] == "P" else 0)
                var_obs = tk.StringVar(value=alum["obs"])
                
                self.asistencia_vars[sec_nombre].append({
                    "presente": var_p, "obs": var_obs, "alumno_ref": alum
                })

                chk = ttk.Checkbutton(frame_tab, text="Presente", variable=var_p)
                chk.grid(row=i+1, column=0, padx=5, pady=2, sticky=tk.W)

                lbl_nombre = ttk.Label(frame_tab, text=f"{alum['nombre']} ({alum['esp']})")
                lbl_nombre.grid(row=i+1, column=1, padx=5, pady=2, sticky=tk.W)

                entry_obs = ttk.Entry(frame_tab, textvariable=var_obs, width=40, state=tk.DISABLED)
                entry_obs.grid(row=i+1, column=2, padx=5, pady=2, sticky=tk.W)

                chk.config(command=lambda p=var_p, e=entry_obs: self.toggle_observacion(p, e))

        frame_bottom = ttk.Frame(self.root)
        frame_bottom.pack(fill=tk.X, padx=10, pady=10)

        btn_generar = ttk.Button(frame_bottom, text="GENERAR PARTE EN WORD", command=self.generar_word)
        btn_generar.pack(side=tk.RIGHT, ipadx=20, ipady=10)

    def toggle_observacion(self, var_presente, entry_widget):
        if var_presente.get() == 1:
            entry_widget.config(state=tk.DISABLED)
            entry_widget.delete(0, tk.END)
        else:
            entry_widget.config(state=tk.NORMAL)
            entry_widget.focus()

    def generar_word(self):
        try:
            self.datos["config"]["MES"] = self.mes_var.get()
            self.datos["config"]["TURNO"] = self.turno_var.get()
            self.datos["config"]["OFICIAL_SEMANA"] = self.oficial_var.get()
            self.datos["config"]["SUBOFICIAL_SEMANA"] = self.suboficial_var.get()

            for sec_nombre, vars_lista in self.asistencia_vars.items():
                for item in vars_lista:
                    alum = item["alumno_ref"]
                    esta_presente = item["presente"].get()
                    alum["estado"] = "P" if esta_presente else "A"
                    alum["obs"] = item["obs"].get() if not esta_presente else ""

            archivo_salida = self.datos["config"]["ARCHIVO_SALIDA"]
            salida, resumen = generar_parte_word(self.datos, archivo_salida)

            msg = (
                f"¡Parte generado exitosamente en tu carpeta!\n\n"
                f"Fuerza Efectiva: {resumen['FE']:02d}\n"
                f"Presentes: {resumen['P']:02d}\n"
                f"Ausentes: {resumen['A']:02d}"
            )
            messagebox.showinfo("Éxito", msg)
            
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al generar el archivo:\n{str(e)}")


if __name__ == "__main__":
    # Lee el Excel desde afuera si existe (por si querés cambiar nombres de alumnos luego),
    # de lo contrario carga el que está empaquetado adentro del EXE.
    ruta_excel_externa = os.path.join(get_output_dir(), "datos_parte.xlsx")
    ruta_excel_interna = get_resource_path("datos_parte.xlsx")

    if os.path.exists(ruta_excel_externa):
        datos_base = cargar_desde_excel(ruta_excel_externa)
    elif os.path.exists(ruta_excel_interna):
        datos_base = cargar_desde_excel(ruta_excel_interna)
    else:
        datos_base = copy.deepcopy(DATOS_POR_DEFECTO)

    root = tk.Tk()
    style = ttk.Style(root)
    if "vista" in style.theme_names():
        style.theme_use("vista")
        
    app = AppPartes(root, datos_base)
    root.mainloop()